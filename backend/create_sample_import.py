#!/usr/bin/env python3
"""
Generate a sample Excel file for bulk product import testing.
Run this script to create sample_products.xlsx in the backend directory.
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# Sample product data
sample_products = [
    {
        "Category": "Gift Boxes",
        "Name": "Premium Gift Box - Small",
        "Slug": "premium-gift-box-small",
        "SKU": "GIFT-BOX-001",
        "HSN Code": "4819",
        "Price": "150.00",
        "Compare At Price": "199.00",
        "GST Percentage": "18.00",
        "Stock": "100",
        "Height CM": "10.5",
        "Width CM": "15.0",
        "Weight G": "200",
        "Is Active": "Yes",
        "Is Bestseller": "Yes",
        "Is Customizable": "No",
        "Is Returnable": "Yes",
        "Bulk Threshold": "50",
        "Min Qty": "1",
        "Is_live": "",  # Empty = import this row
        "Notes": "Popular small gift box",
    },
    {
        "Category": "Gift Boxes",
        "Name": "Premium Gift Box - Medium",
        "Slug": "premium-gift-box-medium",
        "SKU": "GIFT-BOX-002",
        "HSN Code": "4819",
        "Price": "250.00",
        "Compare At Price": "349.00",
        "GST Percentage": "18.00",
        "Stock": "75",
        "Height CM": "15.0",
        "Width CM": "20.0",
        "Weight G": "350",
        "Is Active": "Yes",
        "Is Bestseller": "No",
        "Is Customizable": "Yes",
        "Is Returnable": "Yes",
        "Bulk Threshold": "30",
        "Min Qty": "1",
        "Is_live": "",
        "Notes": "Customizable medium box",
    },
    {
        "Category": "Packaging",
        "Name": "Tissue Paper Pack",
        "Slug": "tissue-paper-pack",
        "SKU": "TISSUE-PAPER-001",
        "HSN Code": "4802",
        "Price": "50.00",
        "Compare At Price": "",
        "GST Percentage": "5.00",
        "Stock": "500",
        "Height CM": "",
        "Width CM": "",
        "Weight G": "100",
        "Is Active": "Yes",
        "Is Bestseller": "Yes",
        "Is Customizable": "No",
        "Is Returnable": "No",
        "Bulk Threshold": "100",
        "Min Qty": "5",
        "Is_live": "",
        "Notes": "Assorted colors",
    },
    {
        "Category": "Ribbon & String",
        "Name": "Silk Ribbon - Gold",
        "Slug": "silk-ribbon-gold",
        "SKU": "RIBBON-GOLD-001",
        "HSN Code": "5807",
        "Price": "75.00",
        "Compare At Price": "99.00",
        "GST Percentage": "12.00",
        "Stock": "200",
        "Height CM": "",
        "Width CM": "2.5",
        "Weight G": "50",
        "Is Active": "Yes",
        "Is Bestseller": "No",
        "Is Customizable": "No",
        "Is Returnable": "Yes",
        "Bulk Threshold": "",
        "Min Qty": "1",
        "Is_live": "",
        "Notes": "5 meter roll",
    },
]

def create_sample_file():
    """Create a sample Excel file for import."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Products"
    
    # Headers
    headers = [
        "Category", "Name", "Slug", "SKU", "HSN Code", "Price",
        "Compare At Price", "GST Percentage", "Stock", "Height CM",
        "Width CM", "Weight G", "Is Active", "Is Bestseller",
        "Is Customizable", "Is Returnable", "Bulk Threshold",
        "Min Qty", "Is_live", "Notes"
    ]
    
    # Write headers with formatting
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.value = header
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # Write sample data
    for row_num, product in enumerate(sample_products, 2):
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=row_num, column=col_num)
            cell.value = product.get(header, "")
            cell.alignment = Alignment(horizontal="left", vertical="center")
    
    # Adjust column widths
    column_widths = {
        "A": 18,  # Category
        "B": 25,  # Name
        "C": 22,  # Slug
        "D": 15,  # SKU
        "E": 12,  # HSN Code
        "F": 12,  # Price
        "G": 15,  # Compare At Price
        "H": 15,  # GST Percentage
        "I": 10,  # Stock
        "J": 12,  # Height CM
        "K": 12,  # Width CM
        "L": 12,  # Weight G
        "M": 12,  # Is Active
        "N": 14,  # Is Bestseller
        "O": 16,  # Is Customizable
        "P": 14,  # Is Returnable
        "Q": 15,  # Bulk Threshold
        "R": 10,  # Min Qty
        "S": 12,  # Is_live
        "T": 20,  # Notes
    }
    
    for col_letter, width in column_widths.items():
        ws.column_dimensions[col_letter].width = width
    
    # Save file
    output_file = "sample_products.xlsx"
    wb.save(output_file)
    print(f"✓ Sample Excel file created: {output_file}")
    print(f"  - Contains {len(sample_products)} sample products")
    print(f"  - Ready to upload to Bulk Import feature")
    print(f"\nUsage:")
    print(f"  1. Go to Django Admin > Catalog > Products")
    print(f"  2. Click 'Bulk Import Products' button")
    print(f"  3. Upload {output_file}")
    print(f"  4. Review the import summary")

if __name__ == "__main__":
    create_sample_file()
