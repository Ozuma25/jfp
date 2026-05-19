#!/usr/bin/env python3
"""
Convert CSV to Excel for bulk product import testing.
Run this script once openpyxl is installed via requirements.txt
"""

import csv
import os

def csv_to_excel():
    """Convert sample_products.csv to sample_products.xlsx"""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        print("ERROR: openpyxl not installed")
        print("Please install dependencies first:")
        print("  cd /home/vishal/Projects/jai_fancy_packs/jfp/backend")
        print("  pip install -r requirements.txt")
        return False
    
    csv_file = "sample_products.csv"
    xlsx_file = "sample_products.xlsx"
    
    if not os.path.exists(csv_file):
        print(f"ERROR: {csv_file} not found")
        return False
    
    # Create workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Products"
    
    # Read CSV and write to Excel
    with open(csv_file, 'r') as f:
        reader = csv.reader(f)
        for row_idx, row in enumerate(reader, 1):
            for col_idx, value in enumerate(row, 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.value = value
                
                # Format header row
                if row_idx == 1:
                    cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
    
    # Adjust column widths
    widths = [18, 25, 22, 15, 12, 12, 15, 15, 10, 12, 12, 12, 12, 14, 16, 14, 15, 10, 12, 20]
    for col_idx, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + col_idx)].width = width
    
    # Save
    wb.save(xlsx_file)
    print(f"✓ Created: {xlsx_file}")
    print(f"  - Ready to upload to Bulk Import feature")
    print(f"\nNext steps:")
    print(f"  1. Go to Django Admin > Catalog > Products")
    print(f"  2. Click 'Bulk Import Products' button")
    print(f"  3. Upload {xlsx_file}")
    return True

if __name__ == "__main__":
    success = csv_to_excel()
    exit(0 if success else 1)
